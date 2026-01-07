
import os
import sys
import time
import json
import datetime
import unicodedata
from urllib.parse import urljoin
from typing import List, Tuple, Dict

import requests
from bs4 import BeautifulSoup
import smtplib
from email.mime.application import MIMEApplication
from email.mime.text import MIMEText
from email.mime.multipart import MIMEMultipart

# PowerPoint
from pptx import Presentation
from pptx.util import Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

# Ημερομηνίες/Ώρες
from dateutil import parser as dateparser
from zoneinfo import ZoneInfo  # stdlib (Python ≥3.9)

# --------------------------------
# ΡΥΘΜΙΣΕΙΣ / ΠΗΓΕΣ
# --------------------------------
SITES = [
    # ---- Core ελληνικά ενεργειακά sites ----
    "https://www.energypress.gr/",
    "https://www.naftemporiki.gr/energy/",
    "https://www.ot.gr/category/energeia/",
    "https://www.capital.gr/energia",
    "https://www.moneyreview.gr/category/business/energy/",
    "https://www.imerisia.gr/energeia",
    "https://www.cnn.gr/oikonomia/energeia",
    "https://www.kathimerini.gr/economy/energy/",
    "https://www.powergame.gr/category/energeia/",
    "https://www.euro2day.gr/news/economy/energeia",
    "https://www.liberal.gr/energeia",
    "https://www.energyin.gr/",
    "https://ecotec.gr/category/energy/",
    "https://www.greenagenda.gr/category/energy/",
    "https://energynews.gr/",

    # ---- Διεθνή sites που καλύπτουν Ελλάδα ----
    "https://greekreporter.com/greek-news/energy/",
    "https://balkangreenenergynews.com/country/greece/",
    "https://energynews.oedigital.com/greece",
    "https://www.reuters.com/world/europe/",
    "https://www.bloomberg.com/energy",
    "https://www.spglobal.com/commodityinsights/en",
    "https://www.ft.com/stream/6d801c42-fd6a-4e71-af32-d398e90e9b5d",

    # ---- ΑΠΕ / Τεχνική ενημέρωση ----
    "https://www.pv-magazine.com/tag/greece/",
    "https://renewablesnow.com/topic/greece/",
    "https://energymag.gr/",
    "https://www.euractiv.com/section/energy/",
    "https://www.euronews.com/tag/energy",

    # ---- Οργανισμοί / Αρχές ----
    "https://ypen.gov.gr/category/anakoinoseis/",
    "https://www.admie.gr/en/news",
    "https://www.raae.gr/anakoinoseis/"
]

# --------------------------------
# Κανονικοποίηση κειμένου
# --------------------------------
ATHENS_TZ = ZoneInfo("Europe/Athens")

def normalize(text: str) -> str:
    """Lowercase + αφαίρεση ελληνικών τόνων/διακριτικών."""
    if not text:
        return ""
    t = text.lower()
    t = ''.join(c for c in unicodedata.normalize('NFD', t) if unicodedata.category(c) != 'Mn')
    return t

# --------------------------------
# Keywords / Negative / Scoring
# --------------------------------
KEYWORDS = [
    # Γενική ενέργεια & αγορά
    "ενεργεια", "ενεργειακη αγορα", "ενεργειακο κοστος", "κιλοβατωρα", "kwh", "mwh",
    "τιμολογια ρευματος", "προμηθευτες ρευματος", "χονδρεμπορικη", "χρηματιστηριο ενεργειας",
    "dam", "day ahead",

    # Δίκτυα
    "ηλεκτρ", "δικτυο μεταφορας", "δικτυο διανομης", "αδμηε", "δεδδηε",
    "διασυνδεση", "interconnector", "ευσταθεια δικτυου", "smart grid", "smart meters",

    # Φωτοβολταϊκά / Ηλιακή
    "φωτοβολ", "φβ", "pv", "πανελ", "modules", "inverter", "μετατροπ", "string", "array",
    "net metering", "αυτοκαταναλωση", "zero feed in",

    # Αιολικά
    "αιολικ", "wind", "ανεμογεννητρ", "turbine", "onshore wind", "offshore wind",
    "repowering", "wind farm",

    # Αποθήκευση
    "μπαταρ", "battery", "bess", "soc", "state of charge", "round trip efficiency",
    "li ion", "lfp", "nmc", "flow battery", "αντλησιοταμιευση", "pumped storage",

    # Υδροηλεκτρικά
    "υδροηλεκτρ", "ταμιευτηρ", "francis", "pelton", "υδροηλεκτρικη ισχυς", "υδατινοι ποροι",

    # Φυσικό αέριο & υδρογονάνθρακες
    "φυσικο αεριο", "gas", "lng", "fsru", "αγωγος", "pipeline",
    "eastmed", "tap", "igb", "upstream", "exploration", "κοιτασμα", "πετρελ", "διυλιστ",

    # Θέρμανση / Κτίρια / Αποδοτικότητα
    "εξοικονομηση", "ενεργειακη αναβαθμιση", "heat pump", "αντλια θερμοτητας",
    "district heating", "τηλεθερμανση", "ενεργειακη κλαση", "θερμικη μονωση",

    # Ρύθμιση / Πολιτική / Αγορά
    "υπεν", "ypen", "ρααευ", "ρυθμιστικη αρχη", "δημοπρασιες απε", "auctions",
    "fit", "cfd", "ppa", "ets",

    # Υδρογόνο & νέα τεχνολογία
    "πρασινο υδρογονο", "green hydrogen", "electrolyzer", "ηλεκτρολυτ",
    "fuel cell", "κυψελη καυσιμου", "power to x", "αποανθρακοποιηση"
]

NEGATIVE = [
    "αθλη", "πολιτισ", "ψυχαγωγ", "μαγειρ", "συνταγ", "μοδα",
    "καιρος", "υγεια", "πανδημ", "κορονο", "τουρισ", "αυτοκινητ",
    "sports", "entertainment"
]

WEIGHTS = {"pv": 3, "bess": 3, "wind": 2, "gas": 2, "policy": 1}
GROUPS = {
    "pv": ["φωτοβολ", "φβ", "pv", "πανελ", "inverter", "net metering", "αυτοκαταναλωση", "zero feed in"],
    "bess": ["μπαταρ", "battery", "bess", "soc", "state of charge", "round trip efficiency", "lfp", "nmc", "flow battery"],
    "wind": ["αιολικ", "wind", "ανεμογεννητρ", "turbine", "onshore wind", "offshore wind", "repowering", "wind farm"],
    "gas": ["φυσικο αεριο", "gas", "lng", "fsru", "αγωγος", "pipeline", "eastmed", "tap", "igb"],
    "policy": ["υπεν", "ypen", "ρααευ", "ρυθμιστικη αρχη", "ppa", "cfd", "fit", "auctions", "ets"]
}

def is_energy_title(title: str) -> bool:
    t = normalize(title)
    if any(n in t for n in NEGATIVE):
        return False
    return any(k in t for k in KEYWORDS)

def score_title(title: str) -> int:
    t = normalize(title)
    score = 0
    for group, words in GROUPS.items():
        if any(w in t for w in words):
            score += WEIGHTS.get(group, 0)
    if any(w in t for w in ["ενεργεια", "ηλεκτρ"]):
        score += 1
    return score

# --------------------------------
# HTTP / Scraping
# --------------------------------
HEADERS = {
    "User-Agent": "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/119.0 Safari/537.36"
}

def absolutize(base: str, href: str) -> str:
    try:
        return urljoin(base, href)
    except Exception:
        return href

def is_probably_article_text(text: str) -> bool:
    """Φιλτράρει πολύ μικρά/άσχετα κείμενα (π.χ. 'Read more', 'Share')."""
    t = normalize(text)
    if len(t) < 8:
        return False
    bad_fragments = ["read more", "περισσοτερα", "share", "mailto:", "javascript:"]
    if any(b in t for b in bad_fragments):
        return False
    return True

def scrape_site(url: str, timeout: int = 12) -> List[Tuple[str, str]]:
    """Επιστρέφει λίστα (title, link) από ένα site."""
    results: List[Tuple[str, str]] = []
    try:
        r = requests.get(url, headers=HEADERS, timeout=timeout)
        r.raise_for_status()
        soup = BeautifulSoup(r.text, "html.parser")
        for a in soup.find_all("a", href=True):
            title = a.get_text(strip=True)
            href = a["href"]
            if not title or not href:
                continue
            if not is_probably_article_text(title):
                continue
            link = absolutize(url, href)
            if link.startswith("#"):
                continue
            results.append((title, link))
        # dedup
        dedup = []
        seen = set()
        for t, l in results:
            key = (normalize(t), normalize(l))
            if key in seen:
                continue
            seen.add(key)
            dedup.append((t, l))
        return dedup
    except Exception as e:
        print(f"[WARN] Σφάλμα στο {url}: {e}", file=sys.stderr)
        return []

# --------------------------------
# Ημερομηνία/Ώρα Δημοσίευσης
# --------------------------------
def try_parse_dt(value: str):
    """Parse οποιαδήποτε ημερομηνία → Europe/Athens. Αν δεν έχει tz → υποθέτουμε UTC."""
    if not value:
        return None
    try:
        dt = dateparser.parse(value)
        if not dt:
            return None
        if not dt.tzinfo:
            dt = dt.replace(tzinfo=ZoneInfo("UTC"))
        return dt.astimezone(ATHENS_TZ)
    except Exception:
        return None

def extract_published_dt_from_article(html: str):
    soup = BeautifulSoup(html, "html.parser")

    candidates = []
    # 1) Meta tags
    for attr in ("property", "name"):
        for key in ("article:published_time", "og:updated_time", "pubdate", "publishdate", "timestamp", "dc.date", "dc.date.issued", "date"):
            tag = soup.find("meta", {attr: key})
            if tag and tag.get("content"):
                candidates.append(tag.get("content"))

    # 2) <time> elements
    for time_tag in soup.find_all("time"):
        dtv = time_tag.get("datetime") or time_tag.get_text(strip=True)
        if dtv:
            candidates.append(dtv)

    # 3) JSON-LD
    for script in soup.find_all("script", {"type": "application/ld+json"}):
        try:
            data = json.loads(script.string or "")
            objs = data if isinstance(data, list) else [data]
            for obj in objs:
                if isinstance(obj, dict):
                    for key in ("datePublished", "dateCreated", "dateModified", "uploadDate"):
                        if obj.get(key):
                            candidates.append(obj[key])
        except Exception:
            pass

    for raw in candidates:
        dt = try_parse_dt(raw)
        if dt:
            return dt
    return None

def fetch_article_published_dt(url: str, timeout: int = 10):
    """Κάνει 1 επιπλέον fetch στη σελίδα άρθρου για να εξάγει ημερομηνία/ώρα."""
    try:
        r = requests.get(url, headers=HEADERS, timeout=timeout)
        r.raise_for_status()
        return extract_published_dt_from_article(r.text)
    except Exception:
        return None

# --------------------------------
# Φίλτρο & Ταξινόμηση (με ημερομηνία)
# --------------------------------
def filter_energy_news(news_list: List[Tuple[str, str]]) -> List[Tuple[str, str, datetime.datetime]]:
    """
    Επιστρέφει [(title, link, published_dt_athens)],
    ταξινομημένα κατά published_dt (desc) και score (desc),
    με dedup και όριο στα extra fetch ανά site για απόδοση.
    """
    seen_links = set()
    tmp: List[Tuple[str, str]] = []

    for title, link in news_list:
        if not title or not link:
            continue
        if is_energy_title(title):
            ln = normalize(link)
            if ln in seen_links:
                continue
            seen_links.add(ln)
            tmp.append((title, link))

    # ταξινόμηση πρώτα κατά score
    tmp.sort(key=lambda x: score_title(x[0]), reverse=True)

    # εμπλουτισμός με published_dt για τα πιο σχετικά
    MAX_ARTICLES_FETCH_DT = 12  # top-N ανά site για ημερομηνία/ώρα
    enriched: List[Tuple[str, str, datetime.datetime]] = []

    for i, (title, link) in enumerate(tmp):
        published_dt = None
        if i < MAX_ARTICLES_FETCH_DT:
            published_dt = fetch_article_published_dt(link)
        enriched.append((title, link, published_dt))

    def sort_key(item):
        title, link, dt = item
        # None → πολύ παλιά, score βοηθά στη σειρά
        base_dt = dt or datetime.datetime.min.replace(tzinfo=ZoneInfo("UTC"))
        return (base_dt, score_title(title))

    enriched.sort(key=sort_key, reverse=True)
    return enriched

# --------------------------------
# Report (text για logs/email body)
# --------------------------------
def fmt_dt(dt: datetime.datetime):
    return dt.strftime("%Y-%m-%d %H:%M") if dt else "—"

def build_text_report(news_dict: Dict[str, List[Tuple[str, str, datetime.datetime]]]) -> str:
    today = datetime.datetime.now(tz=ATHENS_TZ).strftime("%d/%m/%Y %H:%M")
    lines = [f"🔋 DAILY ENERGY NEWS REPORT — {today} (Europe/Athens)", ""]
    for site, items in news_dict.items():
        lines.append(f"============================")
        lines.append(f"SITE: {site}")
        lines.append(f"============================")
        lines.append("")
        if not items:
            lines.append("❗ Δεν βρέθηκαν ενεργειακές ειδήσεις.")
            lines.append("")
        else:
            for title, link, published_dt in items:
                score = score_title(title)
                lines.append(f"• ({score}) [{fmt_dt(published_dt)}] {title}")
                lines.append(f"  {link}")
                lines.append("")
    return "\n".join(lines)

# --------------------------------
# Δημιουργία PowerPoint (.pptx)
# --------------------------------
def add_title_slide(prs: Presentation, title_text: str, subtitle_text: str) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[0])  # Title
    slide.shapes.title.text = title_text
    slide.placeholders[1].text = subtitle_text

def add_bullet_slide(prs: Presentation, heading: str, bullets: List[str]) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[1])  # Title & Content
    slide.shapes.title.text = heading
    tf = slide.shapes.placeholders[1].text_frame
    tf.clear()
    for i, text in enumerate(bullets):
        p = tf.add_paragraph() if i > 0 else tf.paragraphs[0]
        p.text = text
        p.level = 0

def build_pptx(news_dict: Dict[str, List[Tuple[str, str, datetime.datetime]]], output_path: str):
    prs = Presentation()
    today = datetime.datetime.now(tz=ATHENS_TZ).strftime("%d/%m/%Y %H:%M")
    add_title_slide(prs, "Daily Energy News Report", f"Ημερομηνία/Ώρα: {today} (Europe/Athens)")

    empty_total = True
    for site, items in news_dict.items():
        bullets = []
        if not items:
            bullets.append("❗ Δεν βρέθηκαν ενεργειακές ειδήσεις.")
        else:
            empty_total = False
            for title, link, published_dt in items[:10]:  # top 10 ανά site
                score = score_title(title)
                bullets.append(f"[{fmt_dt(published_dt)}] ({score}) {title}\n{link}")
        add_bullet_slide(prs, site, bullets)

    prs.save(output_path)
    return output_path, empty_total

# --------------------------------
# Email με συνημμένο PPTX
# --------------------------------
def send_email_with_attachment(subject: str, body_text: str, attachment_path: str) -> None:
    EMAIL_USERNAME = os.getenv("EMAIL_USERNAME")
    EMAIL_PASSWORD = os.getenv("EMAIL_PASSWORD")
    EMAIL_RECIPIENT = os.getenv("EMAIL_RECIPIENT")
    SMTP_SERVER = os.getenv("SMTP_SERVER", "smtp.gmail.com")
    SMTP_PORT = int(os.getenv("SMTP_PORT", "587"))

    if not EMAIL_USERNAME or not EMAIL_PASSWORD or not EMAIL_RECIPIENT:
        raise RuntimeError("Λείπουν τα env vars: EMAIL_USERNAME / EMAIL_PASSWORD / EMAIL_RECIPIENT")

    msg = MIMEMultipart()
    msg["From"] = EMAIL_USERNAME
    msg["To"] = EMAIL_RECIPIENT
    msg["Subject"] = subject

    # απλό κείμενο στο σώμα (μπορεί να γίνει HTML αν θες)
    msg.attach(MIMEText(body_text, "plain", "utf-8"))

    # συνημμένο PPTX
    with open(attachment_path, "rb") as f:
        part = MIMEApplication(f.read(), Name=os.path.basename(attachment_path))
    part['Content-Disposition'] = f'attachment; filename="{os.path.basename(attachment_path)}"'
    msg.attach(part)

    # SMTP
    server = smtplib.SMTP(SMTP_SERVER, SMTP_PORT)
    server.starttls()
    server.login(EMAIL_USERNAME, EMAIL_PASSWORD)
    server.sendmail(EMAIL_USERNAME, EMAIL_RECIPIENT, msg.as_string())
    server.quit()

# --------------------------------
# MAIN
# --------------------------------
def main():
    all_results: Dict[str, List[Tuple[str, str, datetime.datetime]]] = {}

    for site in SITES:
        print(f"[INFO] Scraping {site} ...")
        raw = scrape_site(site)
        filt = filter_energy_news(raw)
        all_results[site] = filt
        time.sleep(0.5)  # ευγένεια προς sites

    # Text report (για logs/email body)
    text_report = build_text_report(all_results)
    print(text_report)

    # PPTX
    pptx_name = "energy_news_report.pptx"
    pptx_path, empty_total = build_pptx(all_results, pptx_name)

    # Στέλνουμε email (με επιλογή για empty)
    SEND_EMPTY = os.getenv("SEND_EMPTY", "false").lower() == "true"
    if SEND_EMPTY or (not empty_total):
        try:
            send_email_with_attachment("Daily Energy News Report (PPTX)", text_report, pptx_path)
            print("📨 Email (με PPTX) εστάλη επιτυχώς!")
        except Exception as e:
            print(f"❌ Failed to send email: {e}", file=sys.stderr)
    else:
        print("ℹ️ Καμία ενεργειακή είδηση — δεν στέλνω email (μπορείς να ενεργοποιήσεις SEND_EMPTY=true).")

if __name__ == "__main__":
    main()
